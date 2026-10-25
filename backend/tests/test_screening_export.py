import io

from openpyxl import load_workbook

from app.services import screening_export as ex


def test_export_audit_has_ids(monkeypatch):
    import asyncio

    async def fake_audit(year, *a, **k):
        return {"status": "ready", "items": [{"corp_code": "00252001", "corp_name": "해원에스티", "stock_code": "000001",
                                              "period": "2025-12", "rcept_no": "20260330000001", "opinion": "의견거절",
                                              "auditor": "정안회계법인", "scope": ""}]}

    async def fake_company(cc):
        return {"jurir_no": "110111-1234567"}

    monkeypatch.setattr(ex.sc, "audit_nonclean", fake_audit)
    monkeypatch.setattr(ex, "get_company", fake_company)
    body, name = asyncio.run(ex.build("audit", 2025))
    ws = load_workbook(io.BytesIO(body))["감사의견 비적정"]
    rows = list(ws.values)
    assert rows[0][:2] == ("DART 고유번호", "법인등록번호")
    assert rows[1][:4] == ("00252001", "1101111234567", "000001", "해원에스티")
    assert name == "감사의견_비적정_2025.xlsx"


def test_export_convertible_has_bond_sheet(monkeypatch):
    import asyncio

    async def fake_cb(year, *a, **k):
        return {"status": "partial", "items": [{"corp_code": "00000001", "corp_name": "A", "stock_code": "1", "source": "발행결정 추정",
                                                "period": "", "cb_amount": 5, "bw_amount": 0, "total": 5, "next_maturity": "20300101", "count": 1,
                                                "bonds": [{"kind": "CB", "issued": "20260101", "maturity": "20300101", "amount": 5, "rate": ""}]}]}

    async def boom(cc):
        raise RuntimeError("x")

    monkeypatch.setattr(ex.sc, "convertible_bonds", fake_cb)
    monkeypatch.setattr(ex, "get_company", boom)
    body, _ = asyncio.run(ex.build("convertible", 2025))
    wb = load_workbook(io.BytesIO(body))
    assert wb["발행 내역(발행결정 기준)"]["A2"].value == "00000001"
    assert wb["CB·BW 잔액"]["B2"].value in (None, "")  # 법인등록번호 조회 실패는 빈칸
