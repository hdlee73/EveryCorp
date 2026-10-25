"""스크리닝 결과(감사의견 비적정·자본잠식·CB·BW)를 엑셀(.xlsx)로 만든다.

식별번호: DART 고유번호(corp_code)는 항상, 법인등록번호는 기업개황(company.json jurir_no)에서 가져온다.
법인등록번호를 못 가져온 행은 비워 둔다(수집된 목록 자체는 영향 없음).
"""
from __future__ import annotations

import asyncio
import io
import logging
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from . import screening_service as sc
from .company_service import get_company

log = logging.getLogger(__name__)

_CRNO_CONCURRENCY = 4


async def _crnos(codes: list[str]) -> dict[str, str]:
    sem = asyncio.Semaphore(_CRNO_CONCURRENCY)
    out: dict[str, str] = {}

    async def one(cc: str) -> None:
        async with sem:
            try:
                out[cc] = ((await get_company(cc)).get("jurir_no") or "").replace("-", "").strip()
            except Exception as e:  # noqa: BLE001  기업개황을 못 받아도 내보내기는 계속
                log.warning("export crno %s failed: %s", cc, e)
                out[cc] = ""

    await asyncio.gather(*(one(cc) for cc in dict.fromkeys(codes)))
    return out


def _sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200]]) if rows else len(str(h))
        ws.column_dimensions[get_column_letter(i)].width = min(max(width * 1.6 + 2, 10), 40)


def _id_cols(it: dict, crno: dict[str, str]) -> list:
    return [it["corp_code"], crno.get(it["corp_code"], ""), it.get("stock_code", ""), it["corp_name"]]


def _ymd(s: str) -> str:
    return f"{s[:4]}-{s[4:6]}-{s[6:8]}" if len(s) >= 8 and s[:8].isdigit() else (s or "")


_ID_HEAD = ["DART 고유번호", "법인등록번호", "종목코드", "회사명"]


async def build(kind: str, year: int, reprt: str = "11011") -> tuple[bytes, str]:
    """(xlsx 바이트, 파일 이름). kind: audit | capital | convertible. 저장된 수집 결과만 쓰고 새로 수집하지 않는다."""
    wb = Workbook()
    ws = wb.active
    extra = None
    if kind == "audit":
        data = await sc.audit_nonclean(year)
        title = f"감사의견_비적정_{year}"
        crno = await _crnos([i["corp_code"] for i in data["items"]])
        ws.title = "감사의견 비적정"
        _sheet(ws, _ID_HEAD + ["결산", "감사의견", "감사인", "구분", "사업보고서 접수번호"],
               [_id_cols(i, crno) + [i["period"], i["opinion"], i["auditor"], i["scope"], i["rcept_no"]] for i in data["items"]])
    elif kind == "capital":
        data = await sc.capital_impairment(year, reprt)
        title = f"자본잠식_{year}_{sc.REPRT_NAMES.get(reprt, reprt)}"
        crno = await _crnos([i["corp_code"] for i in data["items"]])
        ws.title = "자본잠식"
        _sheet(ws, _ID_HEAD + ["결산", "잠식 구분", "자본총계(원)", "자본금(원)", "자산총계(원)", "부채총계(원)", "재무제표"],
               [_id_cols(i, crno) + [i["period"], i["kind"], i["equity"], i["capital"], i["assets"], i["liabilities"], i["fs_div"]]
                for i in data["items"]])
    elif kind == "convertible":
        data = await sc.convertible_bonds(year)
        title = f"CB_BW_잔액_{year}"
        crno = await _crnos([i["corp_code"] for i in data["items"]])
        ws.title = "회차별 미상환 잔액"
        _sheet(ws, _ID_HEAD + ["구분(CB/BW)", "회차", "종류", "발행일", "만기일", "미상환 잔액(권면액, 원)", "기준일", "근거"],
               [_id_cols(i, crno) + [b["kind"], b["tm"], b["name"], _ymd(b["issued"]), _ymd(b["maturity"]), b["amount"],
                                     _ymd(i["base_date"]), b["basis"]] for i in data["items"] for b in i["bonds"]])
        extra = wb.create_sheet("회사별 합계")
        _sheet(extra, _ID_HEAD + ["기준일", "보고서 결산", "CB 잔액(원)", "BW 잔액(원)", "합계(원)", "가까운 만기", "회차 수", "잔액 근거"],
               [_id_cols(i, crno) + [_ymd(i["base_date"]), i["period"], i["cb_amount"], i["bw_amount"], i["total"],
                                     _ymd(i["next_maturity"]), i["count"], i["source"]] for i in data["items"]])
    else:
        raise ValueError(kind)
    note = wb.create_sheet("안내")
    for line in (f"Every Corp 스크리닝 내보내기 · {title}",
                 f"만든 시각 {datetime.now():%Y-%m-%d %H:%M} · 수집 상태 {data['status']} (ready가 아니면 일부 결과)",
                 "법인등록번호는 DART 기업개황 기준이며 조회하지 못하면 비어 있습니다.",
                 "CB·BW 잔액 기준이 '발행결정 추정'이면 발행 당시 권면총액이라 전환·상환이 반영되지 않았습니다. '정기보고서'는 기준일 현재 미상환 사채 표입니다."):
        note.append([line])
    note.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"{title}.xlsx"
